"""Add pipeline output to the gallery: media, beat times, PPTX export, feed entry.

The pipeline leaves each example as a folder of its own:

    <Style>/<id>/final_video.mp4          the narrated clip
    <Style>/<id>/frame_NN.png             the rendered frames
    <Style>/<id>/mapping_log.txt          frame -> renderer stamp -> narration
    <Style>/<id>/_narration_audio/frames_concat.txt
                                          how long each frame is actually held
    Original_Images/<id>.png              the published diagram

and the gallery wants something different: one media folder per example, the
narration as timed beats, a deck to download, and a line in the feed. This
script is the translation between the two.

BEAT TIMES COME FROM `frames_concat.txt`, NOT `mapping_log.txt`. The log stamps
beats on the renderer's frame counter - evenly spaced whatever was said - while
frames_concat holds each frame for exactly as long as its narration audio runs.
Summing those durations gives the second each line is spoken, which is what the
viewer's caption track needs and what every existing entry carries.

    python tools/add_gallery_samples.py <folder>            # add what is new
    python tools/add_gallery_samples.py <folder> --dry-run  # say what it would do
    python tools/add_gallery_samples.py <folder> --force    # rebuild existing ids too
"""
import io
import json
import os
import re
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FEED = os.path.join(ROOT, 'data/gallery/gallery_items.json')
MEDIA = 'assets/gallery/media'
EXPORTS = 'assets/gallery/exports'

# The pipeline's folder names, and what the gallery calls the same style.
STYLES = {
    'Alpha_Masking': ('alpha', 'Alpha Masking'),
    'Colour_Pop': ('colorpop', 'Colour Pop'),
    'Color_Pop': ('colorpop', 'Colour Pop'),
    'Hopping_BBox': ('hopping', 'Hopping Bbox'),
    'Progressive_Reveal': ('progressive', 'Progressive Reveal'),
    'Sliding_Bbox': ('sliding', 'Sliding Bbox'),
}

SOURCE_W = 1000          # the width every source diagram is stored at
THUMB_W = 640            # and the card image
SLIDE_W = 12191695       # 13.33in in EMU, the width the existing decks use


def run(cmd):
    return subprocess.run(cmd, capture_output=True, text=True)


def probe(path, entries):
    out = run(['ffprobe', '-v', 'error', '-show_entries', entries, '-of', 'csv=p=0', path])
    return out.stdout.strip()


def parse_log(path):
    """[(frame file, renderer stamp, narration)] from a mapping log."""
    rows, cur = [], {}
    for line in io.open(path, encoding='utf-8', errors='replace'):
        m = re.match(r'\s*Frame Image\s*:\s*(\S+)', line)
        if m:
            if cur:
                rows.append(cur)
            cur = {'frame': m.group(1), 'stamp': '', 'text': ''}
            continue
        m = re.match(r'\s*Timestamp\s*:\s*(\S+)', line)
        if m and cur:
            cur['stamp'] = m.group(1)
            continue
        m = re.match(r'\s*Narration\s*:\s*(.*)', line)
        if m and cur:
            cur['text'] = ' '.join(m.group(1).split())
    if cur:
        rows.append(cur)
    return [(r['frame'], r['stamp'], r['text']) for r in rows]


def parse_durations(path):
    """How long each frame is held, in order, from an ffmpeg concat list."""
    return [float(m) for m in re.findall(r'^\s*duration\s+([\d.]+)', io.open(
        path, encoding='utf-8', errors='replace').read(), re.M)]


def deck(frames, rows, out_path, style_label, sample_id):
    """One slide per frame, full bleed, with its line in the speaker notes."""
    from pptx import Presentation
    from pptx.util import Emu

    w, h = [int(x) for x in probe(frames[0], 'stream=width,height').split(',')[:2]]
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(round(SLIDE_W * h / w))
    blank = prs.slide_layouts[6]
    for path, (frame, stamp, text) in zip(frames, rows):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        notes = slide.notes_slide.notes_text_frame
        notes.text = text
        tail = notes.add_paragraph()
        tail.text = '[%s · %s · %s · %s]' % (style_label, sample_id, frame, stamp)
    prs.core_properties.title = 'AnimateBanana — %s — %s' % (style_label, sample_id)
    prs.core_properties.author = 'AnimateBanana'
    prs.save(out_path)


def add(folder, style_dir, name, feed, dry, force):
    style_id, style_label = STYLES[style_dir]
    src = os.path.join(folder, style_dir, name)
    # "CVPR_2025_arch01051 -- fast reveal" is that example, with a note attached
    sample_id = re.split(r'\s+-+\s*', name)[0].strip()

    have = {i['id'] for i in feed['items']}
    if sample_id in have and not force:
        return 'already in the feed'

    video = os.path.join(src, 'final_video.mp4')
    log = os.path.join(src, 'mapping_log.txt')
    concat = os.path.join(src, '_narration_audio', 'frames_concat.txt')
    for p in (video, log, concat):
        if not os.path.exists(p):
            return 'missing %s' % os.path.basename(p)

    frames = sorted(f for f in os.listdir(src) if re.match(r'frame_\d+\.png$', f))
    rows = parse_log(log)
    durations = parse_durations(concat)
    if not (len(frames) == len(rows) == len(durations)):
        return 'frames/log/durations disagree (%d/%d/%d)' % (len(frames), len(rows), len(durations))

    # the source diagram: the published figure, or one already in the gallery
    out_dir = os.path.join(ROOT, MEDIA, sample_id)
    original = os.path.join(folder, 'Original_Images', sample_id + '.png')
    kept = os.path.join(out_dir, 'source.png')
    if not os.path.exists(original) and not os.path.exists(kept):
        return 'no source diagram (Original_Images/%s.png)' % sample_id

    if dry:
        return 'would add: %d frames, %.0fs' % (len(frames), sum(durations))

    os.makedirs(out_dir, exist_ok=True)
    shutil.copy2(video, os.path.join(out_dir, 'narrative.mp4'))
    # the poster is the clip's own first frame, at the clip's own size
    run(['ffmpeg', '-y', '-v', 'error', '-i', video, '-frames:v', '1', '-q:v', '4',
         os.path.join(out_dir, 'poster.jpg')])
    if os.path.exists(original):
        run(['ffmpeg', '-y', '-v', 'error', '-i', original,
             '-vf', "scale='min(%d,iw)':-2:flags=lanczos" % SOURCE_W, kept])
    run(['ffmpeg', '-y', '-v', 'error', '-i', kept,
         '-vf', "scale='min(%d,iw)':-2:flags=lanczos" % THUMB_W, os.path.join(out_dir, 'thumb.png')])

    # ONE BEAT PER LINE, NOT PER FRAME. The renderers emit several frames for a
    # single line - a slide of the highlight box, a few frames of transition -
    # and every one of them carries that line again in the log. Read frame by
    # frame, a 79-second example came out as 77 beats of 11 sentences, each
    # repeated as many times as it had frames. So consecutive frames that say
    # the same thing are one beat: it starts when the first of them does, and
    # the deck gets one slide for it, the last frame of the run, which is the
    # state the line leaves on screen.
    runs, at = [], 0.0
    for i, ((frame, stamp, text), d) in enumerate(zip(rows, durations)):
        if runs and text == runs[-1]['text']:
            runs[-1]['frame'] = frames[i]      # the run's latest frame so far
        else:
            runs.append({'t': round(at, 2), 'text': text, 'frame': frames[i], 'stamp': stamp})
        at += d

    deck([os.path.join(src, r['frame']) for r in runs],
         [(r['frame'], r['stamp'], r['text']) for r in runs],
         os.path.join(ROOT, EXPORTS, sample_id + '.pptx'), style_label, sample_id)

    seconds = float(probe(os.path.join(out_dir, 'narrative.mp4'), 'format=duration') or 0)
    sw, sh = [int(x) for x in probe(kept, 'stream=width,height').split(',')[:2]]
    beats = [{'t': r['t'], 'text': r['text']} for r in runs]

    entry = {
        'id': sample_id,
        'styleId': style_id,
        'style': style_label,
        'thumb': '%s/%s/thumb.png' % (MEDIA, sample_id),
        'source': '%s/%s/source.png' % (MEDIA, sample_id),
        'sourceW': sw,
        'sourceH': sh,
        'narrative': {
            'mp4': '%s/%s/narrative.mp4' % (MEDIA, sample_id),
            'poster': '%s/%s/poster.jpg' % (MEDIA, sample_id),
            'seconds': round(seconds, 2),
        },
        'beats': beats,
        'downloads': {
            'mp4': {
                'href': '%s/%s/narrative.mp4' % (MEDIA, sample_id),
                'name': '%s_%s.mp4' % (sample_id, style_id),
                'note': 'The rendered narrative with its spoken narration.',
            },
            'pptx': {
                'href': '%s/%s.pptx' % (EXPORTS, sample_id),
                'name': '%s_%s.pptx' % (sample_id, style_id),
                'note': '%d static frames, narration in the speaker notes.' % len(runs),
            },
            'svg': None,
        },
    }
    feed['items'] = [i for i in feed['items'] if i['id'] != sample_id] + [entry]
    return 'added: %d beats, %.0fs, %d-slide deck' % (len(beats), seconds, len(runs))


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    dry = '--dry-run' in sys.argv
    force = '--force' in sys.argv
    if not args:
        print(__doc__)
        return
    folder = args[0]

    feed = json.load(io.open(FEED, encoding='utf-8'))
    before = len(feed['items'])
    os.makedirs(os.path.join(ROOT, EXPORTS), exist_ok=True)

    for style_dir in sorted(os.listdir(folder)):
        if style_dir not in STYLES:
            continue
        for name in sorted(os.listdir(os.path.join(folder, style_dir))):
            if not os.path.isdir(os.path.join(folder, style_dir, name)):
                continue
            print('%-38s %s' % (name[:37], add(folder, style_dir, name, feed, dry, force)))

    if dry:
        print('\ndry run: nothing written')
        return
    # the feed reads in style order, then by id, so a new example lands beside
    # its own kind rather than at the end
    order = {s: n for n, s in enumerate(x['id'] for x in feed.get('styles', []))}
    feed['items'].sort(key=lambda i: (order.get(i['styleId'], 99), i['id']))
    io.open(FEED, 'w', encoding='utf-8', newline='\n').write(
        json.dumps(feed, indent=1, ensure_ascii=False) + '\n')
    print('\nfeed: %d items (was %d)' % (len(feed['items']), before))


main()
